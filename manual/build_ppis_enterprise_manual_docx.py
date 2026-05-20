# -*- coding: utf-8 -*-
"""
Build PPIS_Enterprise_End_User_manual.docx from PPIS_Enterprise_End_User_Manual.md
and embed screenshots from the Desktop 'PIS ScreenshotsManual' folder.

Markers in Markdown:
  <<<SCREENSHOTS:COVER>>>     — agency branding images at repository root of manual folder
  <<<SCREENSHOTS:Subfolder>>> — all images inside that subfolder (sorted by name)

Default screenshot root: %USERPROFILE%\\Desktop\\PIS ScreenshotsManual
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt

# --- Markdown subset (aligned with manual/build_docx.py) ---
SCREENSHOT_MARK = re.compile(r"^<<<\s*SCREENSHOTS:\s*([^>]+?)\s*>>>$")


def strip_md_links(text: str) -> str:
    return re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)


def simplify_cell(text: str) -> str:
    t = strip_md_links(text)
    return t.replace("**", "").replace("*", "")


def parse_md_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def maybe_md_table_separator(line: str) -> bool:
    s = line.strip().strip("|")
    if "-" not in s and ":" not in s:
        return False
    chunks = [c.strip() for c in s.split("|")]
    return all(set(c) <= set("-:") for c in chunks if c)


def flush_table(lines: list[str], doc: Document) -> None:
    if len(lines) < 2:
        return
    if maybe_md_table_separator(lines[1]):
        header = parse_md_row(lines[0])
        body_rows = lines[2:]
    else:
        header = parse_md_row(lines[0])
        body_rows = lines[1:]
    ncol = len(header)
    if ncol == 0:
        return
    table = doc.add_table(rows=1 + len(body_rows), cols=ncol)
    table.style = "Table Grid"
    for j, cell in enumerate(header):
        table.rows[0].cells[j].text = simplify_cell(cell)
    for i, brow in enumerate(body_rows):
        cells = parse_md_row(brow)
        while len(cells) < ncol:
            cells.append("")
        for j in range(ncol):
            table.rows[i + 1].cells[j].text = simplify_cell(cells[j])


def add_formatted_run(paragraph, text: str) -> None:
    text = strip_md_links(text)
    parts = re.split(r"(\*\*.+?\*\*|\*.+?\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif (
            part.startswith("*")
            and part.endswith("*")
            and not part.startswith("**")
        ):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


COVER_FILENAMES = (
    "doj-logo.png",
    "ppa-seal.png",
    "bagong-pilipinas.png",
    "redeeming-lives-banner.png",
    "iso-bureau-veritas.png",
)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}


def iter_folder_images(folder: Path) -> list[Path]:
    if not folder.is_dir():
        return []
    files = []
    for p in folder.iterdir():
        if p.is_file() and p.suffix.lower() in IMAGE_EXT:
            files.append(p)
    return sorted(files, key=lambda x: x.name.lower())


def add_screenshots_block(doc: Document, shot_root: Path, key: str, figure_start: list[int]) -> None:
    key = key.strip()
    if key.upper() == "COVER":
        p = doc.add_paragraph()
        p.add_run("Official branding and quality marks (as used in report templates).").italic = True
        for name in COVER_FILENAMES:
            path = shot_root / name
            if path.is_file():
                try:
                    doc.add_picture(str(path), width=Inches(2.35))
                except OSError:
                    doc.add_paragraph(f"[Could not embed image: {name}]")
                cap = doc.add_paragraph()
                r = cap.add_run(f"Figure {figure_start[0]} — {path.stem.replace('-', ' ').title()}")
                r.italic = True
                figure_start[0] += 1
        return

    sub = shot_root / key
    imgs = iter_folder_images(sub)
    if not imgs:
        note = doc.add_paragraph()
        note.paragraph_format.left_indent = Inches(0.2)
        add_formatted_run(
            note,
            f"*Screenshot folder empty or missing: «{key}». Capture from the live PPIS site and save this folder under PIS ScreenshotsManual.*",
        )
        return

    for path in imgs:
        try:
            doc.add_picture(str(path), width=Inches(6.3))
        except OSError:
            doc.add_paragraph(f"[Could not embed image: {path.name}]")
        cap = doc.add_paragraph()
        r = cap.add_run(f"Figure {figure_start[0]} — {key}: {path.stem}")
        r.italic = True
        figure_start[0] += 1


def convert(md_path: Path, out_path: Path, shot_root: Path) -> None:
    doc = Document()
    sect = doc.sections[0]
    sect.left_margin = Inches(1)
    sect.right_margin = Inches(1)
    sect.top_margin = Inches(0.85)
    sect.bottom_margin = Inches(0.85)

    figure_counter = [1]
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_fence = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        mshot = SCREENSHOT_MARK.match(stripped)
        if mshot:
            add_screenshots_block(doc, shot_root, mshot.group(1), figure_counter)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines):
            block = [stripped]
            j = i + 1
            while j < len(lines):
                s2 = lines[j].strip()
                if not s2.startswith("|"):
                    break
                block.append(s2)
                j += 1
            flush_table(block, doc)
            i = j
            continue

        if stripped.startswith("```"):
            in_fence = not in_fence
            i += 1
            continue

        if in_fence:
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph("─" * 56)
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            title = strip_md_links(m.group(2).strip())
            title = title.replace("**", "").replace("*", "")
            if level == 1 and title.strip().lower() == "table of contents":
                doc.add_page_break()
            if level == 1:
                h = doc.add_heading(title, level=0)
                h.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            else:
                doc.add_heading(title, level=min(level - 1, 3))
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            add_formatted_run(p, stripped[2:])
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            p = doc.add_paragraph(style="List Number")
            add_formatted_run(p, text)
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_run(p, stripped[2:])
            i += 1
            continue

        p = doc.add_paragraph()
        add_formatted_run(p, stripped)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def default_shot_root() -> Path:
    return Path.home() / "Desktop" / "PIS ScreenshotsManual"


def main() -> None:
    root = Path(__file__).resolve().parent
    md_path = root / "PPIS_Enterprise_End_User_Manual.md"
    out_path = root / "PPIS_Enterprise_End_User_Manual.docx"
    shot_root = default_shot_root()

    if len(sys.argv) >= 2:
        md_path = Path(sys.argv[1])
    if len(sys.argv) >= 3:
        out_path = Path(sys.argv[2])
    if len(sys.argv) >= 4:
        shot_root = Path(sys.argv[3])

    convert(md_path, out_path, shot_root)
    print(f"Wrote {out_path}")
    print(f"Screenshots from {shot_root}")


if __name__ == "__main__":
    main()
