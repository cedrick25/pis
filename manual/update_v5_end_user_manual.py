"""Update the v5 PPIS end-user manual with current screenshots and workflow text."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SRC = Path(r"e:\Downloads\v5_PIS_End_User_Manual_Revision.docx")
BACKUP = Path(r"e:\Downloads\v5_PIS_End_User_Manual_Revision_backup_before_20260910.docx")
OUT = Path(r"e:\Downloads\v5_PIS_End_User_Manual_Revision_20260910.docx")
SHOT_DIR = ROOT / "PIS_End_User_Manual_Screenshots_20260910"
CATALOG = SHOT_DIR / "screenshot_catalog.json"

RED = RGBColor(0xFF, 0x00, 0x00)
SKIP_CAPTIONS = {"otp page"}


def norm(text: str) -> str:
    text = (text or "").lower().replace("\u2014", " ").replace("\u2013", " ").replace("\u2010", " ")
    text = text.replace("pardone", "pardon")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def iter_block_items(doc: Document):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def first_blip(paragraph: Paragraph):
    blips = paragraph._p.xpath(".//a:blip")
    return blips[0] if blips else None


def replace_image_blob(doc: Document, r_id: str, image_path: Path) -> bool:
    rel = doc.part.rels.get(r_id)
    if rel is None or not image_path.exists():
        return False
    rel.target_part._blob = image_path.read_bytes()
    return True


def set_paragraph_text(paragraph: Paragraph, text: str, red: bool = False, italic: bool = False, size: int = 10):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        run = paragraph.runs[0]
        run.text = text
    else:
        run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.italic = italic
    if red:
        run.font.color.rgb = RED
    else:
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def make_heading(doc: Document, text: str, level: int) -> Paragraph:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def make_body(doc: Document, text: str, red: bool = True) -> Paragraph:
    p = doc.add_paragraph()
    p.style = "Normal"
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    if red:
        run.font.color.rgb = RED
    return p


def make_step(doc: Document, number: int, text: str, red: bool = True) -> Paragraph:
    p = doc.add_paragraph()
    p.style = "Normal"
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    run = p.add_run(f"{number}. {text}")
    run.font.name = "Arial"
    run.font.size = Pt(10)
    if red:
        run.font.color.rgb = RED
    return p


def style_caption(paragraph: Paragraph, text: str):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_text(paragraph, text, red=False, italic=True, size=9)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_picture_paragraph(doc: Document, image_path: Path, width=6.4) -> Paragraph:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    return p


def add_file_keys(index: dict[str, Path], image_path: Path, extra_keys: set[str] | None = None):
    if not image_path.exists():
        return
    stem = re.sub(r"^\d+-", "", image_path.stem)
    keys = {norm(stem), *(extra_keys or set())}
    for key in keys:
        if key:
            index[key] = image_path


def build_catalog_index(records: list[dict]) -> dict[str, Path]:
    index: dict[str, Path] = {}
    for rec in records:
        path = SHOT_DIR / rec["filename"]
        if not path.exists():
            continue
        keys = {
            norm(rec.get("name", "")),
            norm(f"{rec.get('name', '')} {rec.get('kind', '')}"),
            norm(f"{rec.get('module', '')} {rec.get('kind', '')}"),
            norm(rec.get("filename", "").replace(".png", "").split("-", 1)[-1]),
        }
        kind = norm(rec.get("kind", ""))
        module = norm(rec.get("module", ""))
        if kind == "list page":
            keys.add(f"{module} list page")
        if kind == "view":
            keys.add(f"{module} view")
            keys.add(f"{module} view page")
        if kind == "update":
            keys.add(f"{module} update")
            keys.add(f"{module} update page")
        if kind == "attachments":
            keys.add(f"{module} attachments")
            keys.add(f"{module} attachment page")
        if "remove" in kind:
            keys.add(f"{module} delete")
            keys.add(f"{module} delete page")
            keys.add(f"{module} remove confirmation")
        if kind == "add record":
            keys.add(f"{module} add record")
            keys.add(f"{module} add document")
        if "expanded" in kind:
            keys.add("main navigation")
        if kind == "page" and module == "login":
            keys.add("login page")
        if "tabs" in kind:
            keys.add("fact sheet tabs")
            keys.add("fact sheet probation client profile tabs")
        if "click client" in norm(rec.get("name", "")):
            keys.add("fact sheet probation click client name")
        if "client view profile" in norm(rec.get("name", "")):
            keys.add("fact sheet probation client view profile")
        if "worksheet print" in kind or "worksheet print" in norm(rec.get("name", "")):
            keys.add("fact sheet probation worksheet print out sample")
        if "psir print" in kind or "psir print" in norm(rec.get("name", "")):
            keys.add("fact sheet probation psir print out sample")
        if "docket list" in kind or "worksheet and psir" in norm(rec.get("name", "")):
            keys.add("fact sheet probation worksheet and psir overview")
        if "docket worksheet" in kind:
            keys.add("fact sheet probation navigate to worksheet")
        if "docket psir" in kind:
            keys.add("fact sheet probation navigate to psir")
        if "worksheet editor" in kind:
            keys.add("fact sheet probation fill out worksheet")
        if "psir editor" in kind:
            keys.add("fact sheet probation fill out psir")
        if "fact sheet button" in kind:
            keys.add("opening the fact sheet from a docket")
        add_file_keys(index, path, keys)
    return index


def score_match(caption: str, index: dict[str, Path]) -> Path | None:
    key = norm(caption)
    if not key or key in SKIP_CAPTIONS:
        return None
    if key in index:
        return index[key]
    aliases = [
        key.replace(" click client name", ""),
        key.replace(" worksheet print out sample", " worksheet and psir overview"),
        key.replace(" psir print out sample", " fill out psir"),
        key.replace(" fact sheet tabs", " client profile tabs"),
        key.replace(" parole and pardon", " parole and pardone"),
        key.replace(" list page", " list page"),
    ]
    for alias in aliases:
        if alias in index:
            return index[alias]
    best = None
    best_score = 0
    cap_parts = set(key.split())
    for cand, path in index.items():
        cand_parts = set(cand.split())
        overlap = len(cap_parts & cand_parts)
        if overlap >= 3 and overlap > best_score and abs(len(cand_parts) - len(cap_parts)) <= 6:
            best = path
            best_score = overlap
    return best if best_score >= 3 else None


def replace_existing_images(doc: Document, index: dict[str, Path]) -> dict:
    blocks = list(iter_block_items(doc))
    replaced = 0
    skipped = 0
    missing = []
    for i, block in enumerate(blocks):
        if not isinstance(block, Paragraph):
            continue
        blip = first_blip(block)
        if blip is None:
            continue
        caption = ""
        for nxt in blocks[i + 1 : i + 4]:
            if isinstance(nxt, Paragraph) and nxt.text.strip() and first_blip(nxt) is None:
                caption = nxt.text.strip()
                break
        if norm(caption) in SKIP_CAPTIONS or not caption:
            skipped += 1
            continue
        image_path = score_match(caption, index)
        r_id = blip.get(qn("r:embed"))
        if image_path and replace_image_blob(doc, r_id, image_path):
            replaced += 1
        else:
            missing.append(caption)
    return {"replaced": replaced, "skipped": skipped, "missing": missing}


def move_element_after(ref_para: Paragraph, element):
    ref_para._p.addnext(element)


def insert_blocks_after(ref_para: Paragraph, created: list[Paragraph]):
    for item in reversed(created):
        move_element_after(ref_para, item._p)


def find_heading(doc: Document, text: str) -> Paragraph | None:
    target = norm(text)
    for para in doc.paragraphs:
        if para.style and para.style.name.startswith("Heading") and norm(para.text) == target:
            return para
    return None


def replace_if_contains(doc: Document, old: str, new: str) -> int:
    changed = 0
    for para in doc.paragraphs:
        original = para.text
        if old not in original:
            continue
        updated = original.replace(old, new)
        if updated != original:
            set_paragraph_text(paragraph=para, text=updated, red=True)
            changed += 1
    return changed


def update_existing_text(doc: Document) -> int:
    replacements = [
        (
            "4. In the Actions column, use View, Update, Attachments, Remove for the selected row.",
            "4. In the Actions column, use View, Update, Attachments, Fact Sheet, or Remove for the selected row.",
        ),
        (
            "The Worksheet and Post-Sentence Investigation Report (PSIR) are prepared for a client from the client profile. Open the Docket List tab on the client view profile to see each docket row with Worksheet and PSIR status, Edit controls, and print buttons.",
            "The Worksheet and Post-Sentence Investigation Report (PSIR) are prepared from the Docket List tab on the client view profile. Each Worksheet and each PSIR belongs to one docket number. When the status is Not Available, click Edit to start a new Worksheet or PSIR for that docket. When a Worksheet or PSIR already exists, click Edit to continue or update it.",
        ),
        (
            "6. Click Edit in the Worksheet column for the docket that needs the worksheet.",
            "6. Click Edit in the Worksheet column for the docket that needs the worksheet. If the status is Not Available, Edit starts a new worksheet for that docket. If a worksheet already exists, Edit opens the saved record so it can be updated.",
        ),
        (
            "5. Click Edit in the PSIR column for the docket that needs the PSIR to be edited.",
            "5. Click Edit in the PSIR column for the docket that needs the PSIR. If the status is Not Available, Edit starts a new PSIR for that docket. If a PSIR already exists, Edit opens the saved record so it can be updated.",
        ),
        (
            "4. In the confirmation dialog, click Save Changes when creating a new worksheet, or Update Changes when updating an existing worksheet.",
            "4. In the confirmation dialog, click Save Changes when creating a new worksheet, or Update Changes when updating an existing worksheet. Use the back arrow at the top of the worksheet page to return to the client fact sheet.",
        ),
        (
            "4. In the confirmation dialog, click Save Changes when creating a new PSIR, or Update Changes when updating an existing PSIR.",
            "4. In the confirmation dialog, click Save Changes when creating a new PSIR, or Update Changes when updating an existing PSIR. Use the back arrow at the top of the PSIR page to return to the client fact sheet.",
        ),
        (
            "You can view or use Add Investigation Document/Report, Add Supervision Document/Report, Add Note/Other Document, and Add Reporting Date when those buttons are shown.",
            "You can view or use Add Investigation Document/Report, Add Supervision Document/Report, and Add Note/Other Document when those buttons are shown.",
        ),
        (
            "4. In the Actions column, use Update, Attachments for the selected row.",
            "4. In the Actions column, use Update, Fact Sheet, or Attachments for the selected row.",
        ),
        (
            "Use this page to manage client records for parole and pardon workflows.",
            "Use this page to find and open parole and pardon client fact sheets. Click the client name, or click Fact Sheet, to open the client view profile.",
        ),
    ]
    changed = 0
    for old, new in replacements:
        changed += replace_if_contains(doc, old, new)

    tab_old = (
        "Use the tabs in the Client View Profile to navigate through the different sections of the client's record."
    )
    for para in list(doc.paragraphs):
        if para.text.strip() != tab_old:
            continue
        set_paragraph_text(
            para,
            "Use the tabs in the Client View Profile to navigate through the different sections of the client's record.",
            red=False,
        )
        # Rewrite the following bullet paragraphs that describe hidden tabs.
        following = []
        started = False
        for other in doc.paragraphs:
            if other._p is para._p:
                started = True
                continue
            if not started:
                continue
            following.append(other)
            if other.style and other.style.name.startswith("Heading"):
                break
            if other.text.strip().startswith("1."):
                break
        bullets = {
            "rehabilitation": None,
            "reporting dates": None,
            "task list": None,
            "investigation": "Investigation — Review and upload investigation documents and reports related to the client. Use Add Investigation Document/Report when a new file must be attached.",
            "supervision": "Supervision — Review and upload supervision documents and reports for the client. Use Add Supervision Document/Report when a new file must be attached.",
            "notes": "Notes/Other Documents — Keep supplementary notes and other supporting documents for the client. Use Add Note/Other Document to upload a file.",
            "docket list": "Docket List — Review dockets linked to the client. Use this tab to create or update the Worksheet and PSIR for a selected docket.",
        }
        for item in following:
            t = norm(item.text)
            if not t or t.startswith("1 "):
                break
            if "rehabilitation" in t or "reporting dates" in t or "task list" in t:
                el = item._p
                parent = el.getparent()
                if parent is not None:
                    parent.remove(el)
            elif t.startswith("investigation"):
                set_paragraph_text(item, bullets["investigation"], red=True)
            elif t.startswith("supervision"):
                set_paragraph_text(item, bullets["supervision"], red=True)
            elif "notes" in t or "other document" in t:
                set_paragraph_text(item, bullets["notes"], red=True)
            elif "docket list" in t:
                set_paragraph_text(item, bullets["docket list"], red=True)
        changed += 1
        break
    return changed


def insert_docket_factsheet_section(doc: Document, index: dict[str, Path]) -> int:
    anchor = find_heading(doc, "5. Docket Access")
    if anchor is None:
        return 0
    created = []

    def add_heading(text, level=2):
        created.append(make_heading(doc, text, level))

    def body(text):
        created.append(make_body(doc, text, red=True))

    def step(n, text):
        created.append(make_step(doc, n, text, red=True))

    def image(keys, caption):
        path = None
        for key in keys:
            path = index.get(norm(key)) or score_match(key, index)
            if path:
                break
        if not path:
            return
        pic = add_picture_paragraph(doc, path)
        created.append(pic)
        cap = make_body(doc, caption, red=False)
        style_caption(cap, caption)
        created.append(cap)

    add_heading("Opening the Fact Sheet from a docket")
    body(
        "From any docket list, the end-user can open the client fact sheet without leaving Docket Access. "
        "Use Fact Sheet in the Actions column of the selected row."
    )
    step(1, "Open a docket list under Docket Access, for example Probation Investigation.")
    step(2, "Locate the record the end-user needs.")
    step(3, "In the Actions column, click Fact Sheet.")
    step(4, "Wait for the client view profile to open. The page shows the client photo area, profile tabs, and tab content below.")
    image(
        [
            "Opening the Fact Sheet from a docket",
            "opening the fact sheet from a docket fact sheet button",
            "Probation Investigation list page",
        ],
        "Docket list — Fact Sheet in the Actions column",
    )
    image(
        [
            "Opening the Fact Sheet from a docket result",
            "Fact Sheet Probation client view profile",
            "cppo fact sheet full access client profile",
        ],
        "Client view profile opened from Fact Sheet on a docket",
    )
    insert_blocks_after(anchor, created)
    return len(created)


def insert_parole_profile_sections(doc: Document, index: dict[str, Path]) -> int:
    insert_before = find_heading(doc, "Fact Sheet Parole and Pardon (Update)")
    if insert_before is None:
        return 0
    created = []

    def heading(text, level=3):
        created.append(make_heading(doc, text, level))

    def body(text):
        created.append(make_body(doc, text, red=True))

    def step(n, text):
        created.append(make_step(doc, n, text, red=True))

    def image(keys, caption):
        path = None
        for key in keys:
            path = index.get(norm(key)) or score_match(key, index)
            if path:
                break
        if not path:
            return
        pic = add_picture_paragraph(doc, path)
        created.append(pic)
        cap = make_body(doc, caption, red=False)
        style_caption(cap, caption)
        created.append(cap)

    heading("Fact Sheet Parole and Pardon (Open Client Profile)")
    body(
        "The parole and pardon fact sheet uses a client view profile. Open it from the list by clicking the client name or Fact Sheet."
    )
    step(1, "On the Fact Sheet Parole and Pardon list page, locate the client the end-user needs to open.")
    step(2, "Click the client name in the Full Name column, or click Fact Sheet in the Actions column.")
    step(3, "Wait for the client view profile to open. The page shows the client photo area, profile tabs, and tab content below.")
    image(
        [
            "Fact Sheet Parole and Pardon client view profile",
            "fact sheet parole and pardon client view profile",
        ],
        "Fact Sheet Parole and Pardon — client view profile",
    )

    heading("Fact Sheet Parole and Pardon (Client Profile Tabs)")
    body("Use the tabs in the Client View Profile to navigate through the different sections of the client's record.")
    body("Investigation — Review and upload investigation documents and reports related to the client. Use Add Investigation Document/Report when a new file must be attached.")
    body("Supervision — Review and upload supervision documents and reports for the client. Use Add Supervision Document/Report when a new file must be attached.")
    body("Notes/Other Documents — Keep supplementary notes and other supporting documents for the client. Use Add Note/Other Document to upload a file.")
    step(1, "On the client view profile, read the tab labels across the top of the page.")
    step(2, "Click the tab that matches the information the end-user needs to review or update.")
    step(3, "Wait for the tab content to load before using any buttons or tables shown in that section.")
    image(
        [
            "Fact Sheet Parole and Pardon client profile tabs",
            "fact sheet parole and pardon client profile tabs",
        ],
        "Fact Sheet Parole and Pardon — client profile tabs",
    )

    heading("Fact Sheet Parole and Pardon (Upload or Take Photo)")
    body(
        "Take Photo, Upload Photo, Upload Fingerprint, and Add document buttons are shown to the officer who created the client profile, a CPPO, or a System Administrator. Other accounts can review the profile. File actions appear as Restricted."
    )
    step(1, "On the client view profile, locate the client photo area on the right side of the page.")
    step(2, "To upload an existing image file, click Upload Photo.")
    step(3, "In the Upload Profile dialog, choose an image file, then click Confirm.")
    step(4, "To capture a new photo with the device camera, click Take Photo.")
    step(5, "In the Capture Camera dialog, click Open cam, allow camera access if the browser asks for permission, then click Capture.")
    step(6, "Review the captured image. Click Retake if the end-user needs another shot, or click Confirm to save the photo to the client profile.")
    step(7, "Wait for the success message. The client photo area refreshes with the uploaded or captured image.")
    image(
        [
            "Fact Sheet Parole and Pardon upload photo",
            "fact sheet parole and pardon upload photo",
        ],
        "Fact Sheet Parole and Pardon — upload photo",
    )
    image(
        [
            "Fact Sheet Parole and Pardon take photo",
            "fact sheet parole and pardon take photo",
        ],
        "Fact Sheet Parole and Pardon — take photo",
    )

    heading("Fact Sheet Parole and Pardon (Upload Fingerprint)")
    step(1, "On the client view profile, click Upload Fingerprint below the photo buttons.")
    step(2, "Wait for the Upload Fingerprint dialog to open.")
    step(3, "Click the fingerprint image for the finger the end-user needs to upload: Right Thumb, Right Index, Right Middle, Right Ring, or Right Little.")
    step(4, "Select the fingerprint image file from the end-user's computer.")
    step(5, "Repeat the finger selection for any additional fingerprints that must be uploaded.")
    step(6, "Click Confirm to save the fingerprint files to the client profile.")
    step(7, "Wait for the success message, then close the dialog when finished.")

    for item in created:
        insert_before._p.addprevious(item._p)
    return len(created)


def insert_worksheet_save_images(doc: Document, index: dict[str, Path]) -> int:
    created = 0
    caption_map = {
        "Fact Sheet Probation — Fill Out Worksheet": [
            "Fact Sheet Probation Worksheet Save Changes",
            "fact sheet probation worksheet save changes",
        ],
        "Fact Sheet Probation — Fill Out PSIR": [
            "Fact Sheet Probation PSIR Save Changes",
            "fact sheet probation psir save changes",
        ],
    }
    blocks = list(iter_block_items(doc))
    for i, block in enumerate(blocks):
        if not isinstance(block, Paragraph):
            continue
        if block.style and str(block.style.name).startswith("Heading"):
            continue
        caption = block.text.strip()
        keys = None
        for label, key_list in caption_map.items():
            if norm(caption) == norm(label):
                keys = key_list
                break
        if not keys:
            continue
        path = None
        for key in keys:
            path = index.get(norm(key)) or score_match(key, index)
            if path:
                break
        if not path:
            continue
        pic = add_picture_paragraph(doc, path)
        cap = make_body(doc, "Confirmation dialog — Save Changes or Update Changes", red=False)
        style_caption(cap, "Confirmation dialog — Save Changes or Update Changes")
        block._p.addnext(cap._p)
        block._p.addnext(pic._p)
        created += 2
    return created


def cleanup_empty_paragraphs(doc: Document):
    for para in list(doc.paragraphs):
        if first_blip(para) is not None:
            continue
        if para.text.strip() == "" and para.style and para.style.name == "Normal":
            # keep spacing paragraphs that are not the emptied tab bullets with leftover empty runs from our wipe
            if para.runs and all((r.text or "") == "" for r in para.runs) and len(para.runs) <= 2:
                p = para._p
                parent = p.getparent()
                if parent is not None:
                    parent.remove(p)


def main():
    if not SRC.exists():
        raise SystemExit(f"Missing source manual: {SRC}")
    if not CATALOG.exists():
        raise SystemExit(f"Missing screenshot catalog: {CATALOG}")

    if not BACKUP.exists():
        shutil.copy2(SRC, BACKUP)

    catalog = json.loads(CATALOG.read_text(encoding="utf-8"))
    index = build_catalog_index(catalog.get("records", []))

    doc = Document(str(SRC))
    image_stats = replace_existing_images(doc, index)
    text_stats = update_existing_text(doc)
    docket_blocks = insert_docket_factsheet_section(doc, index)
    parole_blocks = insert_parole_profile_sections(doc, index)
    save_images = insert_worksheet_save_images(doc, index)
    cleanup_empty_paragraphs(doc)
    doc.save(str(OUT))

    report = {
        "output": str(OUT),
        "backup": str(BACKUP),
        "screenshots_indexed": len(index),
        "images_replaced": image_stats["replaced"],
        "images_skipped": image_stats["skipped"],
        "images_unmatched": image_stats["missing"],
        "text_updates": text_stats,
        "docket_section_blocks": docket_blocks,
        "parole_section_blocks": parole_blocks,
        "save_dialog_images": save_images,
    }
    report_path = SHOT_DIR / "manual_update_report.json"
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
