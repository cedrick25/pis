"""Replace v3 PPIS manual screenshots and add role-based fact sheet workflows."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SRC = Path(r"e:\Downloads\v3_PIS_End_User_Manual_Revision.docx")
BACKUP = Path(r"e:\Downloads\v3_PIS_End_User_Manual_Revision_backup_before_20260907.docx")
OUT = Path(r"e:\Downloads\v3_PIS_End_User_Manual_Revision.docx")
SHOT_DIR = ROOT / "PIS_End_User_Manual_Screenshots_20260907"
FALLBACK_DIRS = [
    ROOT / "PIS_End_User_Manual_Screenshots_Annotated",
    ROOT / "PIS_End_User_Manual_Screenshots",
]
CATALOG = SHOT_DIR / "screenshot_catalog.json"

# Keep UAMS/cover shots that cannot be recaptured from local PPIS.
SKIP_CAPTIONS = {
    "otp page",
}


def norm(text: str) -> str:
    text = (text or "").lower().replace("\u2014", " ").replace("\u2013", " ")
    text = text.replace("pardone", "pardon")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def iter_block_items(doc: Document):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def first_blip(paragraph: Paragraph):
    blips = paragraph._p.xpath(".//a:blip")
    return blips[0] if blips else None


def replace_image_blob(doc: Document, r_id: str, image_path: Path) -> bool:
    rel = doc.part.rels.get(r_id)
    if rel is None or not image_path.exists():
        return False
    rel.target_part._blob = image_path.read_bytes()
    return True


def add_picture_paragraph(doc: Document, image_path: Path, width=6.4) -> Paragraph:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    return p


def style_caption(paragraph: Paragraph, text: str):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.clear()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def make_heading(doc: Document, text: str, level: int) -> Paragraph:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def make_body(doc: Document, text: str) -> Paragraph:
    p = doc.add_paragraph()
    p.style = "Normal"
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    return p


def make_step(doc: Document, number: int, text: str) -> Paragraph:
    p = doc.add_paragraph()
    p.style = "Normal"
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    run = p.add_run(f"{number}. {text}")
    run.font.name = "Arial"
    run.font.size = Pt(10)
    return p


def add_file_keys(index: dict[str, Path], image_path: Path, extra_keys: set[str] | None = None):
    if not image_path.exists():
        return
    stem = image_path.stem
    stem = re.sub(r"^\d+-", "", stem)
    keys = {norm(stem), *(extra_keys or set())}
    for key in keys:
        if key:
            index[key] = image_path


def index_fallback_folders(index: dict[str, Path]):
    for folder in FALLBACK_DIRS:
        if not folder.exists():
            continue
        for image_path in folder.glob("*.png"):
            add_file_keys(index, image_path)


def build_catalog_index(records: list[dict]) -> dict[str, Path]:
    index = {}
    prefer = (
        "cppo",
        "field officer",
        "clerk",
        "client view",
        "click client",
        "upload photo",
        "take photo",
        "fingerprint",
        "worksheet",
        "psir",
        "docket list",
        "client profile",
    )
    for rec in records:
        path = SHOT_DIR / rec["filename"]
        if not path.exists():
            continue
        blob = norm(f"{rec.get('name', '')} {rec.get('kind', '')} {rec.get('module', '')} {rec.get('filename', '')}")
        if rec.get("kind") in {"list page", "page", "expanded menus", "Remove confirmation", "remove confirmation"} and not any(token in blob for token in prefer):
            continue
        keys = {
            norm(rec.get("name", "")),
            norm(f"{rec.get('name', '')} {rec.get('kind', '')}"),
            norm(f"{rec.get('module', '')} {rec.get('kind', '')}"),
            norm(rec.get("filename", "").replace(".png", "").split("-", 1)[-1]),
        }
        kind = norm(rec.get("kind", ""))
        module = norm(rec.get("module", ""))
        if kind == "list page":
            keys.add(f"{module} list page")
        if kind == "view":
            keys.add(f"{module} view")
            keys.add(f"{module} (view)")
        if kind == "update":
            keys.add(f"{module} update")
            keys.add(f"{module} (update)")
        if kind == "attachments":
            keys.add(f"{module} attachments")
            keys.add(f"{module} (attachments)")
        if "remove" in kind:
            keys.add(f"{module} delete")
            keys.add(f"{module} (delete)")
            keys.add(f"{module} remove confirmation")
        if kind == "add record":
            keys.add(f"{module} add record")
            keys.add(f"{module} (add record)")
            keys.add(f"{module} add document")
            keys.add(f"{module} (add document)")
        if "expanded" in kind:
            keys.add("main navigation")
        if kind == "page" and module == "login":
            keys.add("login page")
        add_file_keys(index, path, keys)
    # Fallbacks first, then current captures overwrite matching keys.
    fallback_index: dict[str, Path] = {}
    index_fallback_folders(fallback_index)
    fallback_index.update(index)
    return fallback_index


def score_match(caption: str, index: dict[str, Path]) -> Path | None:
    key = norm(caption)
    if not key or key in SKIP_CAPTIONS:
        return None
    if key in index:
        return index[key]
    # Common caption variants
    aliases = [
        key.replace(" click client name", ""),
        key.replace(" client view profile", " client view profile"),
        key.replace(" worksheet print out sample", " worksheet and psir overview"),
        key.replace(" psir print out sample", " fill out psir"),
        key.replace(" fact sheet tabs", " client profile tabs"),
    ]
    for alias in aliases:
        if alias in index:
            return index[alias]
    best = None
    best_score = 0
    cap_parts = set(key.split())
    for cand, path in index.items():
        cand_parts = set(cand.split())
        overlap = len(cap_parts & cand_parts)
        if overlap >= 3 and overlap > best_score and abs(len(cand_parts) - len(cap_parts)) <= 6:
            best = path
            best_score = overlap
    return best if best_score >= 3 else None


def replace_existing_images(doc: Document, index: dict[str, Path]) -> dict:
    blocks = list(iter_block_items(doc))
    replaced = 0
    skipped = 0
    missing = []
    for i, block in enumerate(blocks):
        if not isinstance(block, Paragraph):
            continue
        blip = first_blip(block)
        if blip is None:
            continue
        caption = ""
        for nxt in blocks[i + 1 : i + 4]:
            if isinstance(nxt, Paragraph) and nxt.text.strip() and first_blip(nxt) is None:
                caption = nxt.text.strip()
                break
        if norm(caption) in SKIP_CAPTIONS or not caption:
            skipped += 1
            continue
        image_path = score_match(caption, index)
        r_id = blip.get(qn("r:embed"))
        if image_path and replace_image_blob(doc, r_id, image_path):
            replaced += 1
        else:
            missing.append(caption)
    return {"replaced": replaced, "skipped": skipped, "missing": missing}


def move_element_before(ref_para: Paragraph, element):
    ref_para._p.addprevious(element)


def insert_role_workflow_section(doc: Document, index: dict[str, Path]) -> int:
    insert_before = None
    for para in doc.paragraphs:
        if para.text.strip() == "Fact Sheet Parole and Pardon" and para.style.name.startswith("Heading"):
            insert_before = para
            break
    if insert_before is None:
        return 0

    created = []

    def heading(text, level=2):
        p = make_heading(doc, text, level)
        created.append(p)

    def body(text):
        p = make_body(doc, text)
        created.append(p)

    def step(n, text):
        p = make_step(doc, n, text)
        created.append(p)

    def image(keys, caption):
        path = None
        for key in keys:
            path = index.get(norm(key)) or score_match(key, index)
            if path:
                break
        if not path:
            body(f"[{caption}]")
            return
        pic = add_picture_paragraph(doc, path)
        created.append(pic)
        cap = make_body(doc, caption)
        style_caption(cap, caption)
        created.append(cap)

    heading("Fact Sheet Role Workflows (CPPO, Field Officer, Clerk/Other)")
    body(
        "The Probation Fact Sheet now applies a role-based workflow. The same client profile "
        "is shown to all authorized users, but photo, fingerprint, document, worksheet, and "
        "PSIR actions change depending on the signed-in role."
    )
    body(
        "Use this section to confirm what each account type can do after opening a client "
        "from Fact Sheet Probation."
    )

    heading("CPPO and other full-access accounts", 3)
    body(
        "Chief Probation and Parole Officers (CPPO) and other full-access roles can use every "
        "fact-sheet function on any client in their office. This group includes CPPO, System "
        "Administrator, and the other full-access office roles configured in PPIS."
    )
    step(1, "Sign in with a CPPO or other full-access account.")
    step(2, "Open Fact Sheet Probation and click the client name.")
    step(3, "Use Take Photo, Upload Photo, and Upload Fingerprint in the client photo area.")
    step(4, "Use Add Investigation Document/Report, Add Supervision Document/Report, Add Note/Other Document, and Add Reporting Date when those buttons are shown.")
    step(5, "Open the Docket List tab. Edit and print Worksheet and PSIR for any docket, including records created by another officer.")
    image(
        [
            "CPPO fact sheet full access client profile",
            "cppo fact sheet full access client profile",
            "Fact Sheet Probation client view profile",
        ],
        "CPPO / full-access account — client profile with photo and document actions",
    )
    image(
        [
            "CPPO fact sheet full access docket list",
            "Fact Sheet Probation Worksheet and PSIR Overview",
        ],
        "CPPO / full-access account — Docket List with Edit and Print for Worksheet and PSIR",
    )

    heading("Field Officer accounts", 3)
    body(
        "Field Officer accounts can maintain the fact sheet, including photos, fingerprints, "
        "and supporting documents. Worksheet and PSIR edit or print is limited to records "
        "the signed-in officer created. If another officer created the worksheet or PSIR, "
        "the Docket List shows Owner only instead of Edit and Print."
    )
    step(1, "Sign in with a Field Officer account.")
    step(2, "Open Fact Sheet Probation and click the client name.")
    step(3, "Use photo, fingerprint, and document upload actions as needed.")
    step(4, "Open the Docket List tab.")
    step(5, "Click Edit or Print only on a Worksheet or PSIR the officer created. If the record was created by another user, the row shows Owner only.")
    step(6, "To start a Worksheet or PSIR that does not yet exist, use Edit. The new record is saved under the signed-in officer.")
    image(
        [
            "Field Officer fact sheet owner scoped client profile",
            "field officer fact sheet owner scoped client profile",
        ],
        "Field Officer account — client profile with fact-sheet actions available",
    )
    image(
        [
            "Field Officer fact sheet owner scoped docket list",
            "field officer fact sheet owner scoped docket list",
        ],
        "Field Officer account — Docket List showing owner-only Worksheet/PSIR limits",
    )

    heading("Clerk and other accounts", 3)
    body(
        "Clerk and other accounts that are not CPPO/full-access or Field Officer can open "
        "the client fact sheet for viewing only. Photo, fingerprint, add-document, worksheet, "
        "and PSIR actions are hidden or shown as Restricted."
    )
    step(1, "Sign in with a Clerk or other non-officer account.")
    step(2, "Open Fact Sheet Probation and click the client name.")
    step(3, "Review the client profile, tabs, and attached files. Take Photo, Upload Photo, Upload Fingerprint, and Add document buttons are not shown.")
    step(4, "Open Investigation, Supervision, or other file tables. File actions appear as Restricted.")
    step(5, "Open the Docket List tab to review docket status. Worksheet and PSIR Edit/Print controls are not available.")
    image(
        [
            "Clerk other fact sheet restricted client profile",
            "clerk other fact sheet restricted client profile",
        ],
        "Clerk / other account — client profile without photo or add-document buttons",
    )
    image(
        [
            "Clerk other fact sheet restricted docket list",
            "clerk other fact sheet restricted docket list",
        ],
        "Clerk / other account — Docket List without Worksheet/PSIR edit and print",
    )

    heading("What each role can do", 3)
    body(
        "CPPO and other full-access roles: manage photos, fingerprints, documents, worksheets, "
        "and PSIRs for any client record in the office."
    )
    body(
        "Field Officer: manage photos, fingerprints, and documents; create a Worksheet or PSIR "
        "when none exists; edit or print only the Worksheet or PSIR the officer created."
    )
    body(
        "Clerk and other accounts: view the fact sheet and docket information only. The system "
        "does not allow photo capture, file uploads, or Worksheet/PSIR editing from these accounts."
    )

    inserted = 0
    for item in created:
        move_element_before(insert_before, item._p)
        inserted += 1
    return inserted


def update_existing_text(doc: Document) -> int:
    replacements = [
        (
            "The Worksheet and Post-Sentence Investigation Report (PSIR) are prepared for a client from the client profile.",
            "The Worksheet and Post-Sentence Investigation Report (PSIR) are prepared from the client profile. "
            "Availability of Edit and Print depends on the signed-in role. CPPO and other full-access accounts "
            "can manage any record. Field Officers can edit or print only records they created. Clerk and other "
            "accounts can view the Docket List but cannot edit or print Worksheet or PSIR.",
        ),
        (
            "2. To upload an existing image file, click Upload Photo.",
            "2. If the signed-in account is a CPPO, Field Officer, or other allowed role, click Upload Photo. "
            "Clerk and other accounts do not see this button.",
        ),
        (
            "1. On the client view profile, click Upload Fingerprint below the photo buttons.",
            "1. On the client view profile, click Upload Fingerprint below the photo buttons. "
            "This control is shown only to CPPO, Field Officer, and other allowed roles.",
        ),
    ]
    changed = 0
    for para in doc.paragraphs:
        original = para.text
        updated = original
        for old, new in replacements:
            if old in updated:
                updated = updated.replace(old, new)
        if updated != original:
            for run in para.runs:
                run.text = ""
            if para.runs:
                para.runs[0].text = updated
            else:
                para.add_run(updated)
            changed += 1
    return changed


def main():
    if not SRC.exists():
        raise SystemExit(f"Missing source manual: {SRC}")
    if not CATALOG.exists():
        raise SystemExit(f"Missing screenshot catalog: {CATALOG}")

    if not BACKUP.exists():
        shutil.copy2(SRC, BACKUP)

    catalog = {"records": []}
    if CATALOG.exists():
        catalog = json.loads(CATALOG.read_text(encoding="utf-8"))
    index = build_catalog_index(catalog.get("records", []))

    doc = Document(str(SRC))
    image_stats = replace_existing_images(doc, index)
    text_stats = update_existing_text(doc)
    inserted = insert_role_workflow_section(doc, index)
    doc.save(str(OUT))

    report = {
        "output": str(OUT),
        "backup": str(BACKUP),
        "screenshots_indexed": len(index),
        "images_replaced": image_stats["replaced"],
        "images_skipped": image_stats["skipped"],
        "images_unmatched": image_stats["missing"],
        "text_updates": text_stats,
        "role_section_blocks": inserted,
    }
    report_path = SHOT_DIR / "manual_update_report.json"
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
