"""Add Worksheet and PSIR workflows to v2 PIS End User Manual."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SRC = Path(r"c:\Users\Cedrick\Downloads\v2_PIS_End_User_Manual.docx")
OUT_DOWNLOADS = Path(r"c:\Users\Cedrick\Downloads\v2_PIS_End_User_Manual.docx")
OUT_PROJECT = Path(__file__).resolve().parents[1] / "v2_PIS_End_User_Manual.docx"

INSERT_BEFORE = "Fact Sheet Parole and Pardon"


def add_rich_text(paragraph, text, size=10):
    parts = str(text).split("**")
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.bold = idx % 2 == 1


def make_step_paragraph(parent, number, step_text):
    p = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(int(Inches(0.25).twips)))
    ind.set(qn("w:firstLine"), str(int(Inches(-0.25).twips)))
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), str(int(Pt(3).twips)))
    p_pr.append(spacing)
    p.append(p_pr)

    para = Paragraph(p, parent)
    r = para.add_run(f"{number}. ")
    r.font.name = "Arial"
    r.font.size = Pt(10)
    add_rich_text(para, step_text)
    return para


def make_heading_paragraph(parent, text, level=3):
    p = OxmlElement("w:p")
    para = Paragraph(p, parent)
    para.style = f"Heading {level}"
    run = para.add_run(text)
    run.font.name = "Arial"
    return para


def make_normal_paragraph(parent, text=""):
    p = OxmlElement("w:p")
    para = Paragraph(p, parent)
    para.style = "Normal"
    if text:
        add_rich_text(para, text)
    return para


def insert_block_before(ref_para, blocks):
    for block in blocks:
        ref_para._p.addprevious(block._p)


def remove_worksheet_psir_sections(doc):
    """Remove worksheet/PSIR block between Upload Fingerprint caption and Parole section."""
    upload_caption_idx = None
    parole_idx = None
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text.endswith("upload fingerprint"):
            upload_caption_idx = idx
        if text == INSERT_BEFORE and para.style.name == "Heading 2":
            parole_idx = idx
            break
    if upload_caption_idx is None or parole_idx is None or parole_idx <= upload_caption_idx + 1:
        return
    for para in doc.paragraphs[upload_caption_idx + 1 : parole_idx]:
        element = para._element
        element.getparent().remove(element)


def build_blocks(parent):
    blocks = []
    step_no = 0

    def heading(title):
        blocks.append(make_heading_paragraph(parent, title, level=3))

    def intro(text):
        blocks.append(make_normal_paragraph(parent, text))

    def step(text):
        nonlocal step_no
        step_no += 1
        blocks.append(make_step_paragraph(parent, step_no, text))
        return step_no

    def reset_steps():
        nonlocal step_no
        step_no = 0

    def blank():
        blocks.append(make_normal_paragraph(parent))

    heading("Fact Sheet Probation (Worksheet and PSIR Overview)")
    intro(
        "The Worksheet and Post-Sentence Investigation Report (PSIR) are prepared for a "
        "client from the client profile. Open the **Docket List** tab on the client view "
        "profile to see each docket row with **Worksheet** and **PSIR** status, **Edit** "
        "controls, and print buttons."
    )
    blank()

    heading("Fact Sheet Probation (Navigate to Worksheet)")
    reset_steps()
    step("From the left menu, open **Fact Sheet Probation**.")
    step("Locate the client, then click the client name to open the client view profile.")
    step("On the client view profile, click the **Docket List** tab.")
    step(
        "Review the docket table. Each row shows **Docket Number**, **Date Received**, "
        "**Investigating Officer**, **Status**, and the **Worksheet** and **PSIR** columns."
    )
    step(
        "In the **Worksheet** column, note the current status (for example, "
        "**complete**, **incomplete**, or **Not Available**)."
    )
    step(
        "Click **Edit** in the **Worksheet** column for the docket that needs the worksheet."
    )
    step(
        "Wait for the worksheet editor to open on the **Identifying Data** section."
    )
    blank()

    heading("Fact Sheet Probation (Fill Out Worksheet)")
    reset_steps()
    step(
        "On the worksheet screen, use the tabs across the top to move between sections. "
        "Complete each section in order:"
    )
    intro(
        "Identifying Data, Present Offense, Prior Records, Identification Data, "
        "Family Background, Present Situation, Educational History, Employment History, "
        "and Community Background/Environmental Factor."
    )
    step(
        "On the active tab, enter the required information in the fields shown "
        "(for example, **Petitioner's Name**, **Alias/es**, **Date of Initial Interview**, "
        "**Interviewed By**, **Present Address**, and **Permanent Address** on "
        "**Identifying Data**)."
    )
    step("When the current section is complete, click **Save** at the top right of the form.")
    step(
        "In the confirmation dialog, click **Save Changes** when creating a new worksheet, "
        "or **Update Changes** when updating an existing worksheet."
    )
    step(
        "Wait for the success message. The system saves the section and may open the next "
        "worksheet section automatically."
    )
    step(
        "To open another worksheet tab without using **Save**, click the tab name. When the "
        "warning dialog appears, click **Proceed** to open the selected tab."
    )
    step(
        "Repeat data entry and save steps for every worksheet tab until all sections are "
        "filled. The **Worksheet** status on the **Docket List** tab updates to "
        "**complete** when all required sections contain data."
    )
    blank()

    heading("Fact Sheet Probation (Print Worksheet)")
    reset_steps()
    step(
        "Open the client view profile and click the **Docket List** tab, or return to it "
        "after saving worksheet data."
    )
    step("Locate the docket row for the worksheet to print.")
    step("In the **Worksheet** column, click **Print**.")
    step(
        "Wait for the print preview window to open and display the formatted worksheet."
    )
    step(
        "Click **Print** in the preview header. The browser print window opens."
    )
    step(
        "Review the print layout, choose the printer or **Save as PDF**, then click "
        "**Print** in the browser window to finish."
    )
    step(
        "Click the close (**×**) button on the preview window when printing is finished."
    )
    blank()

    heading("Fact Sheet Probation (Navigate to PSIR)")
    reset_steps()
    step("From the left menu, open **Fact Sheet Probation**.")
    step("Click the client name to open the client view profile.")
    step("Click the **Docket List** tab.")
    step(
        "In the **PSIR** column, note the current status (for example, **complete**, "
        "**incomplete**, or **Not Available**)."
    )
    step("Click **Edit** in the **PSIR** column for the docket that needs the PSIR.")
    step(
        "Wait for the PSIR editor to open on the **Identifying Data** section."
    )
    blank()

    heading("Fact Sheet Probation (Fill Out PSIR)")
    reset_steps()
    step(
        "On the PSIR screen, use the tabs across the top to move between sections. "
        "Complete each section in order:"
    )
    intro(
        "Identifying Data, Present Offense, Prior Records, Birth Data and Family Background, "
        "Present Situation, Education and Job History, Medical History, Traits/Characteristics "
        "and Background in the Community, Analysis/Evaluation and Projected Thrust of "
        "Rehabilitation, and Recommendation."
    )
    step(
        "On the active tab, enter the required information in the fields shown "
        "(for example, **Petitioner's Name**, **True Name**, **Source of Information**, "
        "**Alias**, **Height**, **Weight**, **Age**, **Sex**, **Citizenship**, **Religion**, "
        "**Identifying Marks**, **Present Address**, and **Permanent Address** on "
        "**Identifying Data**)."
    )
    step("When the current section is complete, click **Save** at the top right of the form.")
    step(
        "In the confirmation dialog, click **Save Changes** when creating a new PSIR, "
        "or **Update Changes** when updating an existing PSIR."
    )
    step(
        "Wait for the success message. The system saves the section and may open the next "
        "PSIR section automatically."
    )
    step(
        "To open another PSIR tab without using **Save**, click the tab name. When the "
        "warning dialog appears, click **Proceed** to open the selected tab."
    )
    step(
        "Repeat data entry and save steps for every PSIR tab until all required sections "
        "are filled. The **PSIR** status on the **Docket List** tab updates when the "
        "required sections contain data."
    )
    blank()

    heading("Fact Sheet Probation (Print PSIR)")
    reset_steps()
    step(
        "Open the client view profile and click the **Docket List** tab, or return to it "
        "after saving PSIR data."
    )
    step("Locate the docket row for the PSIR to print.")
    step(
        "In the **PSIR** column, click **Short** for the condensed PSIR layout, or "
        "**Long** for the full PSIR that includes the transmittal letter and complete "
        "recommendation blocks."
    )
    step("Wait for the print preview window to open and display the selected PSIR format.")
    step(
        "Click **Print** in the preview header. The browser print window opens."
    )
    step(
        "To save a copy of the PSIR file, click **Download** in the preview header when "
        "shown, then open the downloaded file if needed."
    )
    step(
        "Review the print layout, choose the printer or **Save as PDF**, then click "
        "**Print** in the browser window to finish."
    )
    step(
        "Click the close (**×**) button on the preview window when printing is finished."
    )
    blank()

    return blocks


def find_insert_paragraph(doc):
    for para in doc.paragraphs:
        if para.text.strip() == INSERT_BEFORE and para.style.name == "Heading 2":
            return para
    raise RuntimeError(f"Could not find heading: {INSERT_BEFORE!r}")


def main():
    doc = Document(str(SRC))
    remove_worksheet_psir_sections(doc)
    ref = find_insert_paragraph(doc)
    blocks = build_blocks(ref._parent)
    insert_block_before(ref, blocks)

    saved = []
    for out_path in (OUT_PROJECT, OUT_DOWNLOADS):
        try:
            doc.save(str(out_path))
            saved.append(str(out_path))
        except PermissionError:
            alt = out_path.with_name(out_path.stem + "_worksheet_psir.docx")
            doc.save(str(alt))
            saved.append(str(alt))

    print("Saved:")
    for path in saved:
        print(f"  {path}")
    print(f"Inserted {len(blocks)} paragraphs before {INSERT_BEFORE!r}")


if __name__ == "__main__":
    main()
