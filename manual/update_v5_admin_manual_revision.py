"""Update v5 PPIS manual with fresh admin screenshots (module folders) and light red wording edits.

- Replaces image blobs matched by captions
- Fixes drawing extents so images are not stretched
- Caps display width (~6.5\") and preserves aspect ratio
- Marks wording edits in red, font size 12
- Does NOT rewrite navigation step lists (already finalized)
"""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SRC = Path(r"e:\Downloads\v5_PIS_End_User_Manual_Revision_20260910.docx")
BACKUP = Path(r"e:\Downloads\v5_PIS_End_User_Manual_Revision_20260910_backup_before_admin_rescreenshot.docx")
OUT = Path(r"e:\Downloads\v5_PIS_End_User_Manual_Revision_20260911_admin.docx")
OUT_OVERWRITE = Path(r"e:\Downloads\v5_PIS_End_User_Manual_Revision_20260910.docx")
SHOT_DIR = ROOT / "PIS_End_User_Manual_Screenshots_20260911"
CATALOG = SHOT_DIR / "screenshot_catalog.json"

RED = RGBColor(0xFF, 0x00, 0x00)
SKIP_CAPTIONS = {"otp page"}
MAX_WIDTH_IN = 6.5
EMU_PER_IN = 914400


def norm(text: str) -> str:
    text = (text or "").lower().replace("\u2014", " ").replace("\u2013", " ").replace("\u2010", " ")
    text = text.replace("pardone", "pardon")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def iter_block_items(doc: Document):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def first_blip(paragraph: Paragraph):
    blips = paragraph._p.xpath(".//a:blip")
    return blips[0] if blips else None


def set_run_red_12(run, text: str):
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = RED


def set_paragraph_red_12(paragraph: Paragraph, text: str):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        set_run_red_12(paragraph.runs[0], text)
    else:
        run = paragraph.add_run(text)
        set_run_red_12(run, text)


def image_size_inches(path: Path) -> tuple[float, float]:
    with Image.open(path) as im:
        w_px, h_px = im.size
    # Screenshots were taken at deviceScaleFactor=2 on 1600x900 => logical ~1600 wide
    # Display at max 6.5" wide, preserve aspect ratio (no stretch).
    aspect = h_px / float(w_px) if w_px else 1.0
    width = min(MAX_WIDTH_IN, 6.5)
    height = width * aspect
    # Avoid extremely tall images dominating a page
    if height > 8.5:
        height = 8.5
        width = height / aspect
    return width, height


def set_drawing_size(paragraph: Paragraph, width_in: float, height_in: float):
    cx = int(width_in * EMU_PER_IN)
    cy = int(height_in * EMU_PER_IN)
    for ext in paragraph._p.xpath(".//a:ext"):
        ext.set("cx", str(cx))
        ext.set("cy", str(cy))
    for ext in paragraph._p.xpath(".//wp:extent"):
        ext.set("cx", str(cx))
        ext.set("cy", str(cy))


def replace_image_blob(doc: Document, r_id: str, image_path: Path) -> bool:
    rel = doc.part.rels.get(r_id)
    if rel is None or not image_path.exists():
        return False
    rel.target_part._blob = image_path.read_bytes()
    return True


def add_file_keys(index: dict[str, Path], image_path: Path, extra_keys: set[str] | None = None):
    if not image_path.exists():
        return
    stem = re.sub(r"^\d+-", "", image_path.stem)
    keys = {norm(stem), *(extra_keys or set())}
    for key in keys:
        if key:
            index[key] = image_path


def build_catalog_index(records: list[dict]) -> dict[str, Path]:
    index: dict[str, Path] = {}
    for rec in records:
        path = SHOT_DIR / rec["filename"]
        if not path.exists():
            continue
        keys = {
            norm(rec.get("name", "")),
            norm(f"{rec.get('name', '')} {rec.get('kind', '')}"),
            norm(f"{rec.get('module', '')} {rec.get('kind', '')}"),
            norm(Path(rec.get("filename", "")).stem.split("-", 1)[-1] if "-" in Path(rec.get("filename", "")).stem else Path(rec.get("filename", "")).stem),
        }
        for cap in rec.get("captions") or []:
            keys.add(norm(cap))
        kind = norm(rec.get("kind", ""))
        module = norm(rec.get("module", ""))
        if kind == "list page":
            keys.add(f"{module} list page")
        if kind == "view":
            keys.add(f"{module} view")
            keys.add(f"{module} view page")
            keys.add(f"{module} view")
        if kind == "update":
            keys.add(f"{module} update")
            keys.add(f"{module} update page")
        if kind == "attachments":
            keys.add(f"{module} attachments")
            keys.add(f"{module} attachment page")
        if "remove" in kind:
            keys.add(f"{module} delete")
            keys.add(f"{module} delete page")
            keys.add(f"{module} remove confirmation")
        if kind == "add record":
            keys.add(f"{module} add record")
            keys.add(f"{module} add document")
            keys.add("forms add document")
        if "expanded" in kind:
            keys.add("main navigation")
        if kind == "page" and module == "login":
            keys.add("login page")
        if "tabs" in kind:
            keys.add("fact sheet tabs")
            keys.add("fact sheet probation client profile tabs")
            keys.add("fact sheet parole and pardon client profile tabs")
        if "click client" in norm(rec.get("name", "")):
            keys.add("fact sheet probation click client name")
        if "client view profile" in norm(rec.get("name", "")):
            keys.add("fact sheet probation client view profile")
            keys.add("fact sheet parole and pardon client view profile")
            keys.add("client view profile opened from fact sheet on a docket")
        if "worksheet print" in kind:
            keys.add("fact sheet probation worksheet print out sample")
        if "psir print" in kind:
            keys.add("fact sheet probation psir print out sample")
        if "docket list" in kind or "worksheet and psir" in norm(rec.get("name", "")):
            keys.add("fact sheet probation worksheet and psir overview")
        if "docket worksheet" in kind:
            keys.add("fact sheet probation navigate to worksheet")
        if "docket psir" in kind:
            keys.add("fact sheet probation navigate to psir")
        if "worksheet editor" in kind:
            keys.add("fact sheet probation fill out worksheet")
        if "psir editor" in kind:
            keys.add("fact sheet probation fill out psir")
        if "save dialog" in kind or "save changes" in norm(rec.get("name", "")):
            keys.add("confirmation dialog save changes or update changes")
        if "fact sheet button" in kind:
            keys.add("opening the fact sheet from a docket")
            keys.add("docket list fact sheet in the actions column")
        if module == "cppo" and "client profile" in kind:
            keys.add("cppo full access account client profile with photo and document actions")
        if module == "cppo" and "docket list" in kind:
            keys.add("cppo full access account docket list with edit and print for worksheet and psir")
        if module == "field officer" and "client profile" in kind:
            keys.add("field officer account client profile with fact sheet actions available")
        if module == "field officer" and "docket list" in kind:
            keys.add("field officer account docket list showing owner only worksheet psir limits")
        if module == "clerk or other" and "client profile" in kind:
            keys.add("clerk other account client profile without photo or add document buttons")
        if module == "clerk or other" and "docket list" in kind:
            keys.add("clerk other account docket list without worksheet psir edit and print")
        if "upload photo" in kind:
            keys.add(f"{module} upload photo")
        if "take photo" in kind:
            keys.add(f"{module} take photo")
        if "upload fingerprint" in kind:
            keys.add(f"{module} upload fingerprint")
        add_file_keys(index, path, keys)
    return index


def score_match(caption: str, index: dict[str, Path]) -> Path | None:
    key = norm(caption)
    if not key or key in SKIP_CAPTIONS:
        return None
    if key in index:
        return index[key]
    aliases = [
        key.replace(" click client name", ""),
        key.replace(" worksheet print out sample", " worksheet and psir overview"),
        key.replace(" psir print out sample", " fill out psir"),
        key.replace(" fact sheet tabs", " client profile tabs"),
        key.replace(" parole and pardon", " parole and pardone"),
        key.replace(" list page", " list page"),
        key.replace(" view page", " view"),
        key.replace(" update page", " update"),
        key.replace(" attachment page", " attachments"),
        key.replace(" delete page", " delete"),
    ]
    for alias in aliases:
        if alias in index:
            return index[alias]
    best = None
    best_score = 0
    cap_parts = set(key.split())
    for cand, path in index.items():
        cand_parts = set(cand.split())
        overlap = len(cap_parts & cand_parts)
        if overlap >= 3 and overlap > best_score and abs(len(cand_parts) - len(cap_parts)) <= 6:
            best = path
            best_score = overlap
    return best if best_score >= 3 else None


def replace_existing_images(doc: Document, index: dict[str, Path]) -> dict:
    blocks = list(iter_block_items(doc))
    replaced = 0
    skipped = 0
    resized = 0
    missing = []
    for i, block in enumerate(blocks):
        if not isinstance(block, Paragraph):
            continue
        blip = first_blip(block)
        if blip is None:
            continue
        caption = ""
        for nxt in blocks[i + 1 : i + 4]:
            if isinstance(nxt, Paragraph) and nxt.text.strip() and first_blip(nxt) is None:
                caption = nxt.text.strip()
                break
        # Cover / title image has no usable caption — leave alone
        if not caption or norm(caption) in SKIP_CAPTIONS:
            skipped += 1
            continue
        image_path = score_match(caption, index)
        r_id = blip.get(qn("r:embed"))
        if image_path and replace_image_blob(doc, r_id, image_path):
            width_in, height_in = image_size_inches(image_path)
            set_drawing_size(block, width_in, height_in)
            replaced += 1
            resized += 1
        else:
            missing.append(caption)
    return {"replaced": replaced, "skipped": skipped, "resized": resized, "missing": missing}


def apply_wording_edits(doc: Document) -> list[dict]:
    """Light redundancy cleanup on descriptive prose only. Navigation numbered steps are left alone."""
    edits = []

    # Exact full-paragraph replacements (descriptive intros / redundant module blurbs)
    replacements = [
        (
            "Within the system, users can create, update, and delete records depending on the selected module and performed action. It also provides data export options, enabling users to download displayed table data when needed.",
            "Users can create, update, and delete records by module, and export table data when that option is available.",
        ),
        (
            "For data retrieval and review, users can quickly filter records using inline column search fields or dropdown filter options located above table columns. Results are updated dynamically as filter criteria are entered.",
            "Filter records with the inline column search fields or dropdown filters above the table. Results update as criteria are entered.",
        ),
        (
            "Use this page to search, review, update, attach documents to, or remove probation investigation dockets.",
            "Use this page to search, review, update, attach files to, or remove probation investigation dockets.",
        ),
        (
            "Use this page to manage courtesy investigation dockets under the probation workflow.",
            "Use this page to manage probation courtesy investigation dockets.",
        ),
        (
            "Use this page to manage client records for parole and pardon workflows.",
            "Use this page to find and open parole and pardon client fact sheets.",
        ),
        (
            "Forms – Views and adds documents or forms when permitted.",
            "Forms – View and add documents or forms when permitted.",
        ),
        (
            "Fact Sheet - Maintains client investigation and supervision case records. This serves as the eCase Folder of the client",
            "Fact Sheet – Maintains client investigation and supervision records (eCase Folder).",
        ),
        (
            "Docket Access – Shows the docket records of probation, parole, and pardon",
            "Docket Access – Shows probation, parole, and pardon docket records.",
        ),
        (
            "My Organization - Maintains users, roles, field offices, regions, and permissions.",
            "My Organization – Maintains users, roles, field offices, regions, and permissions.",
        ),
        (
            "From any docket list, the end-user can open the client fact sheet without leaving Docket Access. Use Fact Sheet in the Actions column of the selected row.",
            "From any docket list, open the client fact sheet with Fact Sheet in the Actions column.",
        ),
        (
            "The Worksheet and Post-Sentence Investigation Report (PSIR) are prepared from the Docket List tab on the client view profile. Each Worksheet and each PSIR belongs to one docket number. When the status is Not Available, click Edit to start a new Worksheet or PSIR for that docket. When a Worksheet or PSIR already exists, click Edit to continue or update it.",
            "Prepare the Worksheet and PSIR from the Docket List tab. Each belongs to one docket. If status is Not Available, Edit starts a new record; otherwise Edit opens the existing one.",
        ),
        (
            "Take Photo, Upload Photo, Upload Fingerprint, and Add document buttons are shown to the officer who created the client profile, a CPPO, or a System Administrator. Other accounts can review the profile. File actions appear as Restricted.",
            "Photo, fingerprint, and add-document actions are available to the profile owner, a CPPO, or a System Administrator. Other accounts can review the profile; restricted actions are marked Restricted.",
        ),
    ]

    for para in doc.paragraphs:
        original = para.text.strip()
        if not original:
            continue
        # Skip numbered navigation/action steps entirely
        if re.match(r"^\d+\.\s", original):
            continue
        for old, new in replacements:
            if original == old.strip() or original == old:
                set_paragraph_red_12(para, new)
                edits.append({"from": old, "to": new})
                break
            # Allow soft match ignoring dash variants
            if norm(original) == norm(old) and original != new:
                set_paragraph_red_12(para, new)
                edits.append({"from": original, "to": new})
                break

    # Soften repeated long filter blurb in module intros only when it is a standalone non-step paragraph
    filter_old = (
        "To filter the displayed records, enter search criteria in the inline search fields within the column headers "
        "or choose a value from the available drop-down filters. The table automatically refreshes to show matching results."
    )
    filter_new = (
        "Use the column search fields or drop-down filters to narrow the list; the table refreshes automatically."
    )
    for para in doc.paragraphs:
        t = para.text.strip()
        if re.match(r"^\d+\.\s", t):
            continue
        if norm(t) == norm(filter_old):
            set_paragraph_red_12(para, filter_new)
            edits.append({"from": t, "to": filter_new})

    return edits


def main():
    if not SRC.exists():
        raise SystemExit(f"Missing source manual: {SRC}")
    if not CATALOG.exists():
        raise SystemExit(f"Missing screenshot catalog: {CATALOG}")

    if not BACKUP.exists():
        shutil.copy2(SRC, BACKUP)

    catalog = json.loads(CATALOG.read_text(encoding="utf-8"))
    index = build_catalog_index(catalog.get("records", []))

    doc = Document(str(SRC))
    image_stats = replace_existing_images(doc, index)
    wording_edits = apply_wording_edits(doc)
    doc.save(str(OUT))
    # Also refresh the working revision the user pointed at (backup already saved).
    doc.save(str(OUT_OVERWRITE))

    report = {
        "output": str(OUT),
        "also_updated": str(OUT_OVERWRITE),
        "backup": str(BACKUP),
        "screenshot_dir": str(SHOT_DIR),
        "screenshots_indexed": len(index),
        "images_replaced": image_stats["replaced"],
        "images_skipped": image_stats["skipped"],
        "images_resized": image_stats["resized"],
        "images_unmatched": image_stats["missing"],
        "wording_edits": wording_edits,
        "notes": [
            "User Accounts Update/Delete screenshots could not be captured: UAM http://localhost:8088/user returns HTTP 500, so the admin User Accounts table is empty.",
            "Navigation numbered steps were left unchanged per request.",
            "Image display size capped at ~6.5 inches wide with preserved aspect ratio (no stretch).",
            "No annotation boxes were drawn on screenshots.",
        ],
    }
    report_path = SHOT_DIR / "manual_update_report.json"
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
