import json
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SCREENSHOT_DIR = ROOT / "PIS_End_User_Manual_Screenshots_Annotated"
CATALOG_PATH = SCREENSHOT_DIR / "screenshot_catalog.json"
OUT_PATH = ROOT / "PIS_End_User_Manual.docx"
TEXT_OUT_PATH = ROOT / "manual" / "PIS_End_User_Manual_Draft.md"
FALLBACK_OUT_PATH = ROOT / "PIS_End_User_Manual_updated.docx"
BASE_URL = "http://localhost/pis/"

SUBMITTED_BY = [
    "Submitted by:",
    "Project Management Team",
    "Appcentric Solutions Inc.",
    "Parole and Probation Administration",
]

GROUPS = [
    ("5. Docket Access", [
        "Probation Investigation",
        "Probation Courtesy Investigation",
        "Probation Supervision",
        "Probation Courtesy Supervision",
        "Parole and Pardon Investigation",
        "Parole and Pardon Courtesy Investigation",
        "Parole and Pardon Supervision",
        "Parole and Pardon Courtesy Supervision",
    ]),
    ("6. Docket Routing", [
        "Docket Routing Probation",
        "Docket Routing Pre-Parole",
        "Docket Routing Parole",
        "Docket Routing Pardone",
        "PDL Routing",
    ]),
    ("7. Sent and Inbox", [
        "Sent Probation",
        "Sent Pre-Parole",
        "Sent Parole",
        "Sent Pardone",
        "Sent PDL",
        "Inbox Probation",
        "Inbox Pre-Parole",
        "Inbox Parole",
        "Inbox Pardone",
        "Inbox PDL",
    ]),
    ("8. Fact Sheet", [
        "Fact Sheet Probation",
        "Fact Sheet Parole and Pardone",
        "Fact Sheet PDL",
    ]),
    ("9. Forms", ["Forms"]),
    ("10. My Organization", ["User Accounts", "User Roles", "Field Office", "Region", "Permission"]),
]

PURPOSES = {
    "Dashboard": "Use the dashboard as the starting point after sign-in and to confirm that the account opened the PIS workspace correctly.",
    "Probation Investigation": "Use this page to find, review, update, attach files to, or remove probation investigation dockets.",
    "Probation Courtesy Investigation": "Use this page to manage courtesy investigation dockets under the probation workflow.",
    "Probation Supervision": "Use this page to manage probation supervision dockets and related record actions.",
    "Probation Courtesy Supervision": "Use this page to manage courtesy supervision dockets under the probation workflow.",
    "Parole and Pardon Investigation": "Use this page to find, review, update, attach files to, or remove parole and pardon investigation dockets.",
    "Parole and Pardon Courtesy Investigation": "Use this page to manage courtesy investigation records for parole and pardon cases.",
    "Parole and Pardon Supervision": "Use this page to manage parole and pardon supervision records.",
    "Parole and Pardon Courtesy Supervision": "Use this page to manage courtesy supervision records for parole and pardon cases.",
    "Docket Routing Probation": "Use this queue to route probation dockets to the proper office or receiving user.",
    "Docket Routing Pre-Parole": "Use this queue to start or monitor pre-parole routing work.",
    "Docket Routing Parole": "Use this queue to route parole dockets.",
    "Docket Routing Pardone": "Use this queue to route pardone-related dockets.",
    "PDL Routing": "Use this page for PDL routing actions available to the signed-in account.",
    "Sent Probation": "Use this page to check probation items that have already been sent.",
    "Sent Pre-Parole": "Use this page to check pre-parole items that have already been sent.",
    "Sent Parole": "Use this page to check parole items that have already been sent.",
    "Sent Pardone": "Use this page to check pardone items that have already been sent.",
    "Sent PDL": "Use this page to check PDL items that have already been sent.",
    "Inbox Probation": "Use this page to review probation items received by the office.",
    "Inbox Pre-Parole": "Use this page to review pre-parole items received by the office.",
    "Inbox Parole": "Use this page to review parole items received by the office.",
    "Inbox Pardone": "Use this page to review pardone items received by the office.",
    "Inbox PDL": "Use this page to review PDL items received by the office.",
    "Fact Sheet Probation": "Use this page to find and open probation client fact sheets.",
    "Fact Sheet Parole and Pardone": "Use this page to manage client records for parole and pardone workflows.",
    "Fact Sheet PDL": "Use this page to manage PDL client records.",
    "Forms": "Use this page to view available forms and add documents when the account has permission.",
    "User Accounts": "Use this page to create and maintain user accounts.",
    "User Roles": "Use this page to create and maintain user roles and update permissions assigned to those roles.",
    "Field Office": "Use this page to maintain field office records.",
    "Region": "Use this page to maintain region records.",
    "Permission": "Use this page to maintain permission records used by the system.",
}

ADD_LABELS = {
    "Forms": "Add Document",
    "User Accounts": "Add User Account",
    "User Roles": "Add Role",
    "Field Office": "Add Field Office",
    "Region": "Add Region",
    "Permission": "Add Permission",
    "Fact Sheet PDL": "Add Client",
    "Fact Sheet Parole and Pardone": "Create New Client",
}

FIELD_LABELS = {
    "form-control form-control-sm": "the Search field",
    "sup docket number": "Docket Number",
    "sup client name": "Client Name",
    "csup docket number": "Docket Number",
    "csup client name": "Client Name",
    "sc cinv docket number": "Docket Number",
    "sc cinv client name": "Client Name",
    "docket number display": "Docket Number",
    "client name display": "Client Name",
    "csup view docket": "Docket Number",
    "csup view client": "Client Name",
}

MODAL_MODULES = {
    "User Accounts",
    "User Roles",
    "Field Office",
    "Region",
    "Permission",
}

TEXT_LINES = []


def log_text(line=""):
    TEXT_LINES.append(line)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_rich_text(paragraph, text, size=10):
    parts = str(text).split("**")
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.bold = idx % 2 == 1


def add_numbered_steps(doc, steps):
    for idx, step in enumerate(steps, start=1):
        log_text(f"{idx}. {step.replace('**', '')}")
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{idx}. ")
        r.font.name = "Arial"
        r.font.size = Pt(10)
        add_rich_text(p, step)


def add_bullets(doc, items):
    for item in items:
        log_text(f"• {item.replace('**', '')}")
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("• ")
        r.font.name = "Arial"
        r.font.size = Pt(10)
        add_rich_text(p, item)


def add_note(doc, text, title="Note"):
    log_text(f"{title}: {text.replace('**', '')}")
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    p = cell.paragraphs[0]
    r = p.add_run(f"{title}: ")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(30, 41, 59)
    add_rich_text(p, text, size=9)
    doc.add_paragraph()


def add_picture(doc, record, caption):
    image_path = SCREENSHOT_DIR / record["filename"]
    if not image_path.exists():
        return
    log_text(caption)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(6.35))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(71, 85, 105)


def clean_label(value):
    return str(value or "").replace("×", "X").replace("_", " ").strip()


def humanize_field(value):
    raw = str(value or "").strip()
    if not raw:
        return ""
    low = raw.lower()
    if low in FIELD_LABELS:
        return FIELD_LABELS[low]
    normalized = clean_label(raw).lower()
    if normalized in FIELD_LABELS:
        return FIELD_LABELS[normalized]
    for key, label in FIELD_LABELS.items():
        if key in low or key in normalized:
            return label
    return clean_label(raw)


def is_user_label(value):
    raw = str(value or "").strip()
    label = humanize_field(raw)
    if not label:
        return False
    low = raw.lower()
    if "tbl" in low or "_length" in low or "display" in low or "select2" in low:
        if label not in {"Docket Number", "Client Name"}:
            return False
    if raw in {"text-input", "Data table", "Select field", "X"}:
        return False
    if label.lower() in {"button", "input", "search/filter field"}:
        return False
    if re.match(r"^sc [a-z_]+$", label.lower()) and label not in {"Docket Number", "Client Name"}:
        return False
    return len(label) <= 90


def values(record, key):
    out = []
    for value in (record.get("controls") or {}).get(key, []):
        if not is_user_label(value):
            continue
        label = humanize_field(value)
        if label not in out:
            out.append(label)
    return out


def bold(label):
    return f"**{label}**"


def bold_join(items, limit=6):
    return ", ".join(bold(item) for item in items[:limit])


def find_record(records_by_module, module, kind):
    for record in records_by_module.get(module, []):
        if record["kind"].lower() == kind.lower():
            return record
    return None


def save_button_label(record):
    raw = (record.get("controls") or {}).get("buttons") or []
    if "Confirm" in raw:
        return "Confirm"
    if "Save" in raw:
        return "Save"
    if "Update" in raw:
        return "Update"
    return "Confirm"


def has_data_table(record):
    return bool(values(record, "tables"))


def grid_feature_steps():
    return [
        "To filter the table instantly, type characters directly into the inline search boxes in the column headers, or select a value from any drop-down filter shown above a column. The table updates automatically as you type or select.",
        "To export the currently displayed table data, click **CSV**, **PDF**, or **EXCEL** above the table when those buttons are shown. The system downloads the visible records in the selected format.",
        "To print the table, click **Print**. The native browser print window opens. Review the layout settings, then click **Print** in that window to complete printing.",
    ]


def find_record_by_kind(records_by_module, module, kind):
    for record in records_by_module.get(module, []):
        if (record.get("kind") or "").lower() == kind.lower():
            return record
    return None


def fact_sheet_probation_list_steps(record):
    steps = [
        f"From the left menu, open {bold('Fact Sheet Probation')}.",
        f"To find a client, enter the known name or case number in {bold('Name or case number')}.",
        f"Click {bold('Run search')} to filter the client list.",
        f"In the {bold('Full Name')} column, click the client name link to open the client view profile.",
    ]
    if has_data_table(record):
        steps.extend(grid_feature_steps())
    return steps


def add_fact_sheet_probation_profile_section(doc, records_by_module):
    module = "Fact Sheet Probation"
    open_record = find_record_by_kind(records_by_module, module, "open client name")
    profile_record = find_record_by_kind(records_by_module, module, "client profile")
    upload_photo = find_record_by_kind(records_by_module, module, "upload photo")
    take_photo = find_record_by_kind(records_by_module, module, "take photo")
    upload_fingerprint = find_record_by_kind(records_by_module, module, "upload fingerprint")
    if not profile_record:
        return

    heading = f"{module} (Open Client Profile)"
    log_text("")
    log_text(heading)
    doc.add_heading(heading, level=3)
    add_numbered_steps(doc, [
        f"On the {bold('Fact Sheet Probation')} list page, locate the client you need.",
        f"Click the client name in the {bold('Full Name')} column.",
        "Wait for the client view profile to open. The page shows the client photo area, profile tabs, and tab content below.",
    ])
    if open_record:
        add_picture(doc, open_record, f"{module} — click client name")
    if profile_record:
        add_picture(doc, profile_record, f"{module} — client view profile")

    heading = f"{module} (Client Profile Tabs)"
    log_text("")
    log_text(heading)
    doc.add_heading(heading, level=3)
    log_text("Each tab in the client view profile is used for a different part of the client record.")
    doc.add_paragraph("Each tab in the client view profile is used for a different part of the client record.")
    add_bullets(doc, [
        f"{bold('Investigation')} — Review and upload investigation documents and reports tied to the client. Use {bold('Add Investigation Document/Report')} when a new file must be attached.",
        f"{bold('Supervision')} — Review and upload supervision documents and reports for the client. Use {bold('Add Supervision Document/Report')} when a new file must be attached.",
        f"{bold('Rehabilitation')} — View therapeutic community and restorative justice activities, including dates, venues, phases, stakeholders, and outcomes.",
        f"{bold('Reporting Dates')} — Review reporting-date photos and records. Use {bold('Add Reporting Date')} when a new reporting entry or photo must be recorded.",
        f"{bold('Notes/Other Documents')} — Store supplementary notes and other supporting documents for the client. Use {bold('Add Note/Other Document')} to upload a file.",
        f"{bold('Task List')} — Review tasks assigned to the client or office, including assignment date, assigned user, and status.",
        f"{bold('Docket List')} — Review dockets linked to the client from investigation, supervision, and related workflows.",
    ])
    add_numbered_steps(doc, [
        "On the client view profile, read the tab labels across the top of the page.",
        "Click the tab that matches the information you need to review or update.",
        "Wait for the tab content to load before using any buttons or tables shown in that section.",
    ])

    heading = f"{module} (Upload or Take Photo)"
    log_text("")
    log_text(heading)
    doc.add_heading(heading, level=3)
    add_numbered_steps(doc, [
        "On the client view profile, locate the client photo area on the left side of the page.",
        f"To upload an existing image file, click {bold('Upload Photo')}.",
        f"In the {bold('Upload Profile')} dialog, choose an image file, then click {bold('Confirm')}.",
        f"To capture a new photo with the device camera, click {bold('Take Photo')}.",
        f"In the {bold('Capture Camera')} dialog, click {bold('Open cam')}, allow camera access if the browser asks for permission, then click {bold('Capture')}.",
        "Review the captured image. Click **Retake** if you need another shot, or click **Confirm** to save the photo to the client profile.",
        "Wait for the success message. The client photo area refreshes with the uploaded or captured image.",
    ])
    if upload_photo:
        add_picture(doc, upload_photo, f"{module} — upload photo")
    if take_photo:
        add_picture(doc, take_photo, f"{module} — take photo")

    heading = f"{module} (Upload Fingerprint)"
    log_text("")
    log_text(heading)
    doc.add_heading(heading, level=3)
    add_numbered_steps(doc, [
        "On the client view profile, click **Upload Fingerprint** below the photo buttons.",
        f"Wait for the {bold('Upload Fingerprint')} dialog to open.",
        "Click the fingerprint image for the finger you need to upload: **Right Thumb**, **Right Index**, **Right Middle**, **Right Ring**, or **Right Little**.",
        "Select the fingerprint image file from your computer.",
        "Repeat the finger selection for any additional fingerprints that must be uploaded.",
        f"Click {bold('Confirm')} to save the fingerprint files to the client profile.",
        "Wait for the success message, then close the dialog when finished.",
    ])
    if upload_fingerprint:
        add_picture(doc, upload_fingerprint, f"{module} — upload fingerprint")


def permission_list_steps(record):
    steps = [
        f"From the left menu, open {bold('My Organization')}, then click {bold('Permission')}.",
        "Review the **Permission List** table with columns **ID**, **Permission Name**, **Type**, **Detail**, and **Action**.",
        f"To find a permission, type in the {bold('Search')} box at the top right of the list.",
        f"Use {bold('Update')} in the {bold('Action')} column to edit a permission, or click {bold('Add Permission')} to create a new one.",
    ]
    if has_data_table(record):
        steps.extend(grid_feature_steps())
    return steps


def permission_add_steps():
    return [
        f"On the {bold('Permission')} list page, click {bold('Add Permission')} at the top right.",
        f"Wait for the {bold('New Permission')} modal dialog to appear.",
        f"Enter {bold('Permission Name')}, select {bold('Type')} ({bold('VIEW')} or {bold('ACTION')}), and enter {bold('Detail')}.",
        f"The {bold('Detail')} field identifies the page, API, or action controlled by the permission.",
        f"Click {bold('Confirm')} inside the modal dialog to save the new permission.",
        "Wait for the success message and for the modal dialog to close. The new permission appears in the list.",
    ]


def permission_update_steps():
    return [
        f"On the {bold('Permission')} list page, locate the permission record that needs correction.",
        f"Click {bold('Update')} in the {bold('Action')} column.",
        f"Wait for the {bold('Update Permission')} modal dialog to appear.",
        f"Review and update {bold('Permission Name')}, {bold('Type')}, and {bold('Detail')} as needed.",
        f"Click {bold('Confirm')} inside the modal dialog to save the changes.",
        "Wait for the modal dialog to close. The permission list refreshes with the updated record.",
    ]


def list_steps(module, record):
    if module == "Fact Sheet Probation":
        return fact_sheet_probation_list_steps(record)
    if module == "Fact Sheet PDL":
        steps = [
            f"From the left menu, open {bold(module)}.",
            f"If the page has tabs, select the needed tab, such as {bold('Investigation')} or {bold('Supervision')}.",
            f"To find a record, enter the known information in {bold('Search Client')} when shown.",
            f"In the {bold('Actions')} column, use {bold_join(['View', 'Update', 'Attachments'])} for the selected row.",
        ]
        if has_data_table(record):
            steps.extend(grid_feature_steps())
        return steps
    if module == "Permission":
        return permission_list_steps(record)
    fields = values(record, "fields")
    buttons = values(record, "buttons")
    tabs = values(record, "tabs")
    steps = [f"From the left menu, open {bold(module)}."]
    if tabs:
        steps.append(f"If the page has tabs, select the needed tab, such as {bold_join(tabs, 4)}.")
    if fields:
        search_fields = [
            field for field in fields
            if "search" in field.lower() or "name" in field.lower() or "docket" in field.lower() or "number" in field.lower()
        ]
        selected = search_fields or fields
        if len(selected) == 1 and selected[0] == "the Search field":
            steps.append(f"To find a record, enter the known information in {bold('the Search field')}.")
        else:
            steps.append(f"To find a record, enter the known information in {bold_join(selected, 4)}.")
    if "Run search" in buttons:
        steps.append(f"Click {bold('Run search')} to apply the search.")
    elif "Search" in buttons and module == "User Accounts":
        steps.append(f"Click {bold('Search')} to apply the search.")
    actions = [button for button in ["View", "Update", "Attachments", "Remove"] if button in buttons]
    if actions:
        steps.append(f"In the {bold('Actions')} column, use {bold_join(actions)} for the selected row.")
    else:
        steps.append("Review the records shown on the page and follow any visible office procedure for the selected queue.")
    if has_data_table(record):
        steps.extend(grid_feature_steps())
    return steps


def action_title(kind):
    titles = {
        "View": "View",
        "Update": "Update",
        "Attachments": "Attachments",
        "Remove confirmation": "Delete",
        "Add record": "Add Record",
        "create client": "Create Client",
    }
    return titles[kind]


def modal_update_steps(module, record):
    field_text = bold_join(values(record, "fields"), 8) if values(record, "fields") else "the required fields shown in the dialog"
    save = save_button_label(record)
    return [
        f"On the {bold(module)} list page, locate the record that needs correction.",
        f"Click {bold('Update')} in the {bold('Actions')} column.",
        "Wait for the modal dialog box to appear on the screen.",
        f"Review the current details and update {field_text} inside the modal dialog.",
        f"Click {bold(save)} inside the modal dialog to submit the changes.",
        "If the system displays a confirmation prompt, review the information and click **Confirm** to save the updated record.",
        "Wait for the modal dialog to close. The table refreshes and displays the updated record.",
    ]


def page_update_steps(module, record):
    field_text = bold_join(values(record, "fields"), 8) if values(record, "fields") else "the required fields shown on the form"
    save = save_button_label(record)
    return [
        f"On the {bold(module)} list page, locate the record that needs correction.",
        f"Click {bold('Update')} in the {bold('Actions')} column.",
        "Wait for the update screen to open.",
        f"Review the current details and update {field_text}.",
        f"Click {bold(save)} to save after checking the information.",
        "Wait for the system to return to the list page or refresh the record display.",
    ]


def action_steps(kind, record):
    module = record["module"]
    fields = values(record, "fields")
    buttons = values(record, "buttons")

    if kind == "View":
        return [
            f"On the {bold(module)} list page, locate the record you need to review.",
            f"Click {bold('View')} in the {bold('Actions')} column.",
            "Read the details on the view page. Use this screen for checking information before making any changes.",
        ]

    if kind == "Update":
        if module == "Permission":
            return permission_update_steps()
        if module in MODAL_MODULES:
            return modal_update_steps(module, record)
        return page_update_steps(module, record)

    if kind == "Attachments":
        return [
            f"On the {bold(module)} list page, click {bold('Attachments')} for the correct record.",
            "Wait for the attachments screen to open.",
            f"Complete the attachment details, such as {bold('Kind')}, {bold('Type')}, and {bold('Upload File')} when shown.",
            f"Click {bold('Confirm')} to upload the file.",
            f"Use {bold('View')}, {bold('Download')}, or {bold('Delete')} only for the attachment listed in the table.",
        ]

    if kind == "Remove confirmation":
        remove = "Remove permanently" if "Remove permanently" in buttons else "Remove"
        return [
            f"On the {bold(module)} list page, locate the sample, test, or authorized record.",
            f"Click {bold('Remove')} in the {bold('Actions')} column.",
            "Wait for the confirmation modal dialog box to appear.",
            "Read the confirmation message carefully. The system warns that removal cannot be undone.",
            f"Click {bold('Cancel')} inside the modal dialog to stop, or click {bold(remove)} inside the modal dialog only when deletion is authorized.",
            "Wait for the modal dialog to close. The table refreshes to reflect the remaining records.",
        ]

    if kind == "Add record":
        if module == "Permission":
            return permission_add_steps()
        add_label = ADD_LABELS.get(module, "Add")
        field_text = bold_join(fields, 8) if fields else "all required fields"
        save = save_button_label(record)
        if module == "Forms":
            return [
                f"On the {bold('Forms')} list page, click {bold(add_label)}.",
                "Wait for the upload screen to open.",
                f"Enter {bold('Uploader')}, {bold('Kind')}, and select a file in {bold('Upload File')}.",
                f"Click {bold('Confirm')} to save after checking the entry.",
                "Wait for the system to confirm the upload and return to the forms list.",
            ]
        if module in MODAL_MODULES or module == "Fact Sheet PDL":
            return [
                f"On the {bold(module)} page, click {bold(add_label)}.",
                "Wait for the modal dialog box to appear on the screen.",
                f"Enter {field_text} inside the modal dialog.",
                f"Click {bold(save)} inside the modal dialog to submit the new record.",
                "If the system displays a confirmation prompt, review the information and click **Confirm** to save the entry.",
                "Wait for the modal dialog to close. The table refreshes and displays the newly added record.",
            ]
        return [
            f"Click {bold(add_label)}.",
            f"Fill in {field_text}.",
            f"Click {bold(save)} to save after checking the entry.",
        ]

    if kind == "create client":
        field_text = bold_join(values(record, "fields"), 12) if values(record, "fields") else "all required fields"
        save = save_button_label(record)
        return [
            f"Open {bold('Fact Sheet Parole and Pardone')} from the left menu.",
            f"When your role allows client creation, open the {bold('Create New Client')} page.",
            f"Select {bold('Client Type')} as {bold('Parole')} or {bold('Pardone')}.",
            f"Enter {field_text}.",
            f"Click {bold(save)} to save the new client record.",
            "Wait for the success message. The new client appears in the Parole or Pardone tab on the fact sheet list.",
        ]

    return [f"Follow the controls shown on the {bold(kind)} page."]


def add_module_section(doc, module, records_by_module):
    records = records_by_module[module]
    list_record = find_record(records_by_module, module, "list page") or records[0]
    log_text("")
    log_text(module)
    doc.add_heading(module, level=2)
    purpose = PURPOSES.get(module, f"Use this module to complete the {module} workflow.")
    log_text(purpose)
    doc.add_paragraph(purpose)
    add_numbered_steps(doc, list_steps(module, list_record))
    add_picture(doc, list_record, f"{module} list page")

    if module == "Fact Sheet Probation":
        add_fact_sheet_probation_profile_section(doc, records_by_module)

    module_kinds = ["View", "Update", "Attachments", "Remove confirmation", "Add record"]
    if module == "Fact Sheet Parole and Pardone":
        module_kinds.insert(0, "create client")

    for kind in module_kinds:
        record = find_record(records_by_module, module, kind)
        if not record:
            continue
        title = action_title(kind)
        heading = f"{module} ({title})"
        log_text("")
        log_text(heading)
        doc.add_heading(heading, level=3)
        if kind == "Remove confirmation":
            add_note(
                doc,
                f"Use {bold('Remove')} only for sample/test records or records approved for deletion. For documentation, stop at the confirmation dialog unless removal is required.",
                "Deletion warning",
            )
        add_numbered_steps(doc, action_steps(kind, record))
        add_picture(doc, record, heading)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for style_name, size, color in [
        ("Heading 1", 18, "1F2937"),
        ("Heading 2", 14, "334155"),
        ("Heading 3", 11, "475569"),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def add_heading_block(doc, text, level=1):
    log_text("")
    log_text(text)
    doc.add_heading(text, level=level)


def build_doc():
    global TEXT_LINES
    TEXT_LINES = []

    data = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
    records_by_module = defaultdict(list)
    for record in data["records"]:
        records_by_module[record["module"]].append(record)

    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(72)
    r = title.add_run("PIS End User Manual")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor(15, 23, 42)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Probation and Parole Information System")
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(51, 65, 85)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(28)
    r = meta.add_run(
        f"End-User Operating Guide\nWebsite: {BASE_URL}\nDocument Date: {date.today().strftime('%B %d, %Y')}"
    )
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(71, 85, 105)

    log_text("PIS End User Manual")
    log_text("Probation and Parole Information System")
    log_text(f"End-User Operating Guide")
    log_text(f"Website: {BASE_URL}")
    log_text(f"Document Date: {date.today().strftime('%B %d, %Y')}")

    doc.add_page_break()

    add_heading_block(doc, "Table of Contents", level=1)
    add_bullets(doc, [
        "1. What is PIS?",
        "2. System Modules",
        "3. Logging In",
        "4. Navigation",
        "5. Docket Access",
        "6. Docket Routing",
        "7. Sent and Inbox",
        "8. Fact Sheet",
        "9. Forms",
        "10. My Organization",
        "11. Safe Editing, Deleting, and Troubleshooting",
    ])
    doc.add_page_break()

    add_heading_block(doc, "1. What is PIS?", level=1)
    what_is_pis = [
        "PIS stands for Probation and Parole Information System. As the name implies, the system is used for managing probation, parole, and pardon docketing, routing, fact sheets, and organizational records across field offices and workflow queues.",
        "The system automatically processes and aggregates tracking and docketing workflows. Users can search, view, update, attach files to, route, and remove records according to their assigned role and permissions.",
        "Users can add, update, and delete records after selecting the corresponding module and action. They can export visible table data through **CSV**, **PDF**, and **EXCEL** buttons, and print records through the **Print** button where those controls appear on data tables.",
        "Users can filter records instantly by typing directly into inline column search headers or by selecting values from drop-down filters above table columns. Filtered results update automatically as criteria are entered.",
    ]
    for paragraph in what_is_pis:
        log_text(paragraph.replace("**", ""))
        doc.add_paragraph(paragraph.replace("**", ""))

    add_note(
        doc,
        "This manual guides probation and parole staff through the PIS website using the actual screens and labels shown in the system.",
        "Manual purpose",
    )
    add_note(
        doc,
        "Menus may vary by user role. If a module is not visible, ask the system administrator to check your account permissions.",
        "Role access",
    )
    add_note(
        doc,
        "Red boxes and arrows labeled **Click** show the control to select in each screenshot.",
        "Screenshot indicators",
    )

    add_heading_block(doc, "2. System Modules", level=1)
    add_bullets(doc, [
        "**Docket Access** - Work with probation, parole, and pardon docket records.",
        "**Docket Routing** - Route dockets between offices or workflow queues.",
        "**Sent** and **Inbox** - Review items that were sent or received.",
        "**Fact Sheet** - Open and maintain client fact sheet records.",
        "**Forms** - View and add documents when permitted.",
        "**My Organization** - Maintain users, roles, field offices, regions, and permissions.",
    ])

    add_heading_block(doc, "3. Logging In", level=1)
    login_record = records_by_module["Login"][0]
    add_numbered_steps(doc, [
        f"Open {bold(BASE_URL)} in a supported browser.",
        f"Type your username in the {bold('Email address')} field.",
        f"Type your password in the password field.",
        f"Click {bold('SIGN IN')}.",
        "Wait for PIS to open the first page allowed for your account.",
    ])
    add_picture(doc, login_record, "Login page")

    add_heading_block(doc, "4. Navigation", level=1)
    if "Navigation" in records_by_module:
        add_numbered_steps(doc, [
            f"Use the left menu to move between {bold('Docket Access')}, {bold('Docket Routing')}, {bold('Fact Sheet')}, {bold('Forms')}, and {bold('My Organization')}.",
            f"Click a menu group such as {bold('Probation')}, {bold('Sent')}, {bold('Inbox')}, or {bold('My Organization')} to show its submenu.",
            "Click the exact module name you need to open.",
        ])
        add_picture(doc, records_by_module["Navigation"][0], "Main navigation")

    for heading, modules in GROUPS:
        add_heading_block(doc, heading, level=1)
        for module in modules:
            if module in records_by_module:
                add_module_section(doc, module, records_by_module)

    add_heading_block(doc, "11. Safe Editing, Deleting, and Troubleshooting", level=1)
    add_bullets(doc, [
        f"Use {bold('View')} when you only need to check record details.",
        f"Use {bold('Update')} only when you are authorized to change the selected record. Review the fields inside the update screen or modal dialog, then click {bold('Confirm')} or {bold('Update')} to save.",
        f"Use {bold('Attachments')} to upload, view, download, or delete supporting files tied to the selected record.",
        f"Use {bold('Remove')} only for records approved for deletion. For training, use sample/test records and click {bold('Cancel')} inside the confirmation modal unless deletion is required.",
        "To filter a table instantly, type in the inline column search headers or use drop-down filters above the columns.",
        "To export table data, click **CSV**, **PDF**, or **EXCEL** when those buttons appear above the table.",
        "To print table data, click **Print**, review the browser print layout settings, then execute print from the browser print window.",
        f"If search results look wrong, clear the search field and click {bold('Run search')} again.",
        "If a menu is missing, ask the administrator to check your role and permission settings.",
    ])
    add_note(
        doc,
        "The annotated screenshots used in this manual are saved in the **PIS_End_User_Manual_Screenshots_Annotated** folder.",
        "Screenshot folder",
    )

    doc.add_paragraph()
    signoff = doc.add_paragraph()
    signoff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    signoff.paragraph_format.space_before = Pt(24)
    for idx, line in enumerate(SUBMITTED_BY):
        log_text("")
        log_text(line)
        if idx > 0:
            signoff.add_run("\n")
        run = signoff.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.bold = idx == 0

    footer = doc.sections[-1].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("PIS End User Manual")
    r.font.name = "Arial"
    r.font.size = Pt(8)

    TEXT_OUT_PATH.write_text("\n".join(TEXT_LINES) + "\n", encoding="utf-8")

    try:
        doc.save(OUT_PATH)
        print(OUT_PATH)
    except (PermissionError, OSError):
        doc.save(FALLBACK_OUT_PATH)
        print(FALLBACK_OUT_PATH)
    print(TEXT_OUT_PATH)


if __name__ == "__main__":
    build_doc()
